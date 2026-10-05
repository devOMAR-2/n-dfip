import re
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

import os
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'docs', 'N-DFIP-BRD.md')
OUT = os.path.join(ROOT, 'docs', 'N-DFIP-BRD.docx')
PRIMARY = RGBColor(0x91, 0x32, 0x46)
DARK = RGBColor(0x5C, 0x1A, 0x2C)

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2)
sec.top_margin = sec.bottom_margin = Cm(2)

styles = doc.styles
styles['Normal'].font.name = 'Calibri'
styles['Normal'].font.size = Pt(10.5)
for lvl, size in [(1, 16), (2, 13), (3, 11.5), (4, 11)]:
    st = styles[f'Heading {lvl}']
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = PRIMARY if lvl <= 2 else DARK


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)')


def add_inline(par, text, bold=False, size=None, color=None):
    for part in INLINE.split(text):
        if not part:
            continue
        b, i, code = bold, False, False
        if part.startswith('**') and part.endswith('**'):
            part, b = part[2:-2], True
        elif part.startswith('`') and part.endswith('`'):
            part, code = part[1:-1], True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            part, i = part[1:-1], True
        run = par.add_run(part)
        run.bold = b
        run.italic = i
        if code:
            run.font.name = 'Consolas'
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color


# ---------- Title page ----------
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(140)
add_inline(p, 'N-DFIP', bold=True, size=36, color=PRIMARY)
p = doc.add_paragraph()
add_inline(p, 'National Diabetic Foot Intelligence Platform', size=18, color=DARK)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(24)
add_inline(p, 'Business Requirements Document', bold=True, size=20)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(36)
for line in ['Version 0.1 (draft)', 'Date: 5 October 2026', 'Status: Draft for client review', 'Classification: Confidential']:
    q = doc.add_paragraph()
    add_inline(q, line, size=11)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- Body ----------
lines = open(SRC, encoding='utf-8').read().split('\n')
i = 0
skip_title = True


def table_rows(block):
    rows = []
    for ln in block:
        cells = [c.strip() for c in ln.strip().strip('|').split('|')]
        rows.append(cells)
    return rows


while i < len(lines):
    ln = lines[i]
    s = ln.strip()
    if not s or s == '---':
        i += 1
        continue
    if s.startswith('# '):
        i += 1  # title already on the title page
        continue
    if skip_title and s.startswith('**National Diabetic'):
        skip_title = False
        i += 1
        continue
    m = re.match(r'^(#{2,4}) (.*)', s)
    if m:
        lvl = len(m.group(1)) - 1
        doc.add_heading(m.group(2), level=lvl)
        i += 1
        continue
    if s.startswith('|'):
        block = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            block.append(lines[i])
            i += 1
        rows = [r for r in table_rows(block) if not all(re.fullmatch(r':?-{2,}:?', c or '--') for c in r)]
        ncols = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=ncols)
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, r in enumerate(rows):
            for ci in range(ncols):
                cell = t.cell(ri, ci)
                cell.text = ''
                par = cell.paragraphs[0]
                txt = r[ci] if ci < len(r) else ''
                add_inline(par, txt, bold=(ri == 0), size=9, color=(RGBColor(0xFF, 0xFF, 0xFF) if ri == 0 else None))
                if ri == 0:
                    shade(cell, '913246')
        doc.add_paragraph()
        continue
    if s.startswith('>'):
        par = doc.add_paragraph()
        par.paragraph_format.left_indent = Cm(1)
        add_inline(par, s.lstrip('> ').strip(), size=10, color=DARK)
        par.runs[0].italic = True
        i += 1
        continue
    m = re.match(r'^(\s*)- (.*)', ln)
    if m:
        depth = len(m.group(1)) // 2
        par = doc.add_paragraph(style='List Bullet' if depth == 0 else 'List Bullet 2')
        add_inline(par, m.group(2))
        i += 1
        continue
    m = re.match(r'^\d+\. (.*)', s)
    if m:
        par = doc.add_paragraph(style='List Number')
        add_inline(par, m.group(1))
        i += 1
        continue
    par = doc.add_paragraph()
    add_inline(par, s)
    i += 1

# Footer with page numbers
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_inline(footer, 'N-DFIP BRD v0.1 (draft) · Page ', size=8)
run = footer.add_run()
for tag, text in [('begin', None), (None, 'PAGE'), ('end', None)]:
    if tag:
        fc = OxmlElement('w:fldChar')
        fc.set(qn('w:fldCharType'), tag)
        run._r.append(fc)
    else:
        it = OxmlElement('w:instrText')
        it.set(qn('xml:space'), 'preserve')
        it.text = text
        run._r.append(it)
run.font.size = Pt(8)

doc.save(OUT)
print('saved', OUT)
